from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
root=Path(__file__).resolve().parent.parent
path=root/'俚雨_畢業製作企畫書.docx'
d=Document(path)
paras=list(d.paragraphs)
first=next(i for i,p in enumerate(paras) if p.text.startswith('一　'))
for p in paras[:first]: p._element.getparent().remove(p._element)
heading=paras[first]
def before(text,style=None):
 p=heading.insert_paragraph_before(text,style)
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 return p
p=before('作品名稱 俚雨','Title')
p.paragraph_format.space_before=Cm(10.5)
p.paragraph_format.space_after=Cm(5.8)
p=before('組員'); p.paragraph_format.space_after=Pt(5)
for text in ['1121834　徐曼玲','1121819　游祐瑄','1121821　陳品霓','1122032　黃麗文','1122006　黃翊禎']:
 p=before(text); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.1
p=before('指導老師　陳彥彰、趙伯堯'); p.paragraph_format.space_before=Pt(12)
breakp=before('')
sect=OxmlElement('w:sectPr')
for child in d.sections[0]._sectPr:
 if child.tag not in [qn('w:headerReference'),qn('w:footerReference'),qn('w:pgNumType')]:
  import copy
  sect.append(copy.deepcopy(child))
typ=OxmlElement('w:type'); typ.set(qn('w:val'),'nextPage'); sect.append(typ)
breakp._p.get_or_add_pPr().append(sect)
body=d.sections[-1]
body.header.is_linked_to_previous=False
body.footer.is_linked_to_previous=False
body.header.paragraphs[0].text='俚雨'
body.header.paragraphs[0].runs[0].font.size=Pt(10)
body.header_distance=Cm(1)
body.footer_distance=Cm(1.1)
num=OxmlElement('w:pgNumType'); num.set(qn('w:start'),'1'); body._sectPr.append(num)
p=body.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
d.save(path)
