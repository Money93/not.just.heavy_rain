from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

root=Path(__file__).resolve().parent.parent
d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2.2)
s.left_margin=s.right_margin=Cm(2.5)
for st in d.styles:
 if st.type==1:
  st.font.name='Times New Roman'; st.font.size=Pt(12); st.font.color.rgb=RGBColor(0,0,0); st.font.italic=False
  for el in list(st._element.iter(qn('w:pBdr'))): el.getparent().remove(el)
  fonts=st._element.get_or_add_rPr().get_or_add_rFonts()
  for key,val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','標楷體'),('cs','Times New Roman')]: fonts.set(qn('w:'+key),val)
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: fonts.attrib.pop(qn('w:'+key),None)
  st.paragraph_format.space_after=Pt(6)
  st.paragraph_format.line_spacing=1.35
for i in range(1,6):
 st=d.styles[f'Heading {i}']; st.font.size=Pt(16 if i==1 else 12); st.font.bold=True
 st.paragraph_format.space_before=Pt(10 if i<3 else 5)
 st.paragraph_format.keep_with_next=True
 st.paragraph_format.first_line_indent=Pt(0)
 st.paragraph_format.left_indent=Cm((i-1)*0.35)
d.styles['Title'].font.size=Pt(26)
d.styles['Normal'].paragraph_format.first_line_indent=Pt(24)

def p(t): return d.add_paragraph(t)
def h(t,n):
 x=d.add_heading(t,n)
 bm=OxmlElement('w:bookmarkStart'); bm.set(qn('w:id'),str(len(headings))); bm.set(qn('w:name'),'H'+str(len(headings)))
 x._p.insert(0,bm)
 end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(len(headings))); x._p.append(end)
 headings.append((t,n,'H'+str(len(headings))))
 return x
headings=[]
x=d.add_paragraph('俚雨\n畢業製作企畫書','Title'); x.alignment=1; x.paragraph_format.space_before=Cm(4); x.paragraph_format.space_after=Cm(3)
for t in ['115學年度專題製作','作品類型　互動','組員','1121834　徐曼玲','1121819　游祐瑄','1121821　陳品霓','1122032　黃麗文','1122006　黃翊禎','指導老師　陳彥彰、趙伯堯']:
 x=p(t); x.alignment=1; x.paragraph_format.first_line_indent=Pt(0)
d.add_page_break()
x=d.add_paragraph('目錄','Title'); x.alignment=1
toc_anchor=p('')
d.add_page_break()
h('壹、企畫概要與創作背景',1)
h('一、作品概要',2)
p('《俚雨》以世界各國形容暴雨的特殊俚語為創作起點，結合雨傘設計與互動投影，將抽象的語言轉化為趣味的互動體驗。作品規劃將六個國家的暴雨俚語轉化為動態圖像，並設計六款具有不同俚語圖案的雨傘。觀眾透過選傘、撐傘及移動雨傘，與對應國家的俚語雨景互動，體驗不同文化所描繪的雨景。')
h('二、創作動機',2)
p('不同國家對暴雨有著獨特的語言想像，例如英國的「下貓下狗」、希臘的「下椅腳」及巴西的「下蛇與蜥蜴」。這些表達方式提供了將語言轉化為圖像與互動的創作素材。隨著全球氣候變遷，極端降雨及其帶來的影響，已成為各地共同關注的環境議題。本作品希望以有趣的文化體驗吸引觀眾，進一步引導對氣候議題的理解與關注。')
h('三、議題定位與創作目標',2)
h('(一) 永續發展目標',3)
p('本作品對應 SDG 13「氣候行動」及其 13.3「氣候變遷教育與意識提升」。透過暴雨俚語、互動投影，以及互動成果的 QR Code 與社群延伸內容，分享氣候知識及日常可實踐的環境友善行動，將環境關注延伸至日常生活。')
h('(二) 創作目標',3)
h('1. 文化轉譯與互動參與',4)
p('將六國暴雨俚語轉化為動態雨景，以六款實體雨傘串連圖案、語意及操作回饋，讓觀眾認識不同文化的語言特色與多元表達方式。')
h('2. 氣候意識與日常延伸',4)
h('(1) 氣候知識傳達',5)
p('以互動體驗連結極端降雨對生活與環境的影響，提升觀眾對氣候議題的關注。')
h('(2) 環境友善行動',5)
p('以 QR Code 與社群內容提供延伸閱讀及日常行動資訊，鼓勵觀眾持續關心環境。')

d.add_page_break()
h('貳、核心概念與文化內容',1)
h('一、概念轉譯',2)
p('作品以「暴雨俚語—圖像設計—實體雨傘—互動雨景」串連文化內容與體驗。各國俚語中的物件、動物與想像轉化為雨傘圖案及投影動畫，讓觀眾藉由持傘參與，感受不同文化對同一自然現象的表達方式。')
h('二、六國俚語內容規劃',2)
h('(一) 作品說明列舉的文化意象',3)
p('英國「下貓下狗」以貓與狗作為雨景圖像；希臘「下椅腳」以椅腳作為視覺物件；巴西「下蛇與蜥蜴」以蛇與蜥蜴建立場景特色。三者分別發展對應的傘面圖案、動態物件與俚語說明。')
h('(二) 其餘國家與語言查核',3)
p('作品總規模為六個國家；其餘三國由團隊完成資料查核與內容評估後定案。各國俚語的原文、中文意譯、使用地區及語境，應於製作前確認，並建立來源紀錄，避免將直譯意象當成普遍用法。')
h('三、內容呈現與議題連結',2)
h('(一) 展場文化內容',3)
p('投影與作品說明應讓觀眾辨認雨傘所對應的國家、俚語及雨景意象。文字與動畫共同協助理解，並維持各國內容一致的資訊結構與閱讀順序。')
h('(二) 體驗後的氣候內容',3)
p('觀眾完成互動後，透過互動成果的 QR Code 連結延伸內容，閱讀極端降雨、氣候變遷與生活環境的相關資訊。社群內容持續分享氣候知識及可實踐的環境友善行動，延伸展場體驗。')

d.add_page_break()
h('參、視覺設計與互動體驗',1)
h('一、雨傘與動態圖像設計',2)
h('(一) 六款雨傘設計',3)
p('以六國暴雨俚語為靈感，設計六款具有不同俚語圖案的雨傘。傘面圖像需兼顧文化辨識、整體視覺特色與影像辨識需求，使雨傘同時作為展示物件與互動媒介。')
h('(二) 投影與標誌設計',3)
p('將俚語意象發展為趣味動態圖像，統一色彩、物件造形與動畫節奏，並確保觀眾與雨傘不遮蔽主要資訊。LOGO 以「俚雨」為作品名稱，應能應用於傘面、投影、展場說明及社群內容；最終美術風格經場景試作後確認。')
h('二、觀眾互動流程',2)
h('(一) 選傘與進場',3)
p('觀眾閱讀作品與操作說明，從六款雨傘中選擇一款，進入指定互動區。系統辨識所選雨傘，建立與對應國家俚語雨景的連結。')
h('(二) 撐傘與移動互動',3)
p('觀眾撐傘並移動雨傘，系統透過影像辨識取得操作資訊，即時更新投影中的對應雨景。具體回饋方式與有效活動範圍由互動原型測試確認，並以現場引導協助操作。')
h('(三) 成果與延伸閱讀',3)
p('互動結束後呈現互動成果與 QR Code，讓觀眾連結氣候知識及社群延伸內容。成果產生方式、呈現時間及下一位觀眾的重置流程，於整合測試階段確認。')
h('三、展場配置',2)
p('以牆面作為主要投影面，配置雨傘展示與取用區、持傘互動區及成果閱讀區。攝影機需能擷取有效操作範圍，電腦整合辨識與動畫後輸出投影。投影距離、鏡頭高度及觀眾站位依實際場地調整，並保留撐傘與移動所需的安全空間。')

d.add_page_break()
h('肆、技術架構與設備規劃',1)
h('一、系統架構',2)
p('系統以「雨傘及觀眾操作—影像擷取—辨識與位置追蹤—雨景動畫—即時投影—互動成果與 QR Code」為基本流程。開發重點為六款雨傘與六國雨景的正確對應，以及撐傘、移動操作的即時回饋。')
h('二、辨識與畫面整合',2)
h('(一) 雨傘辨識與追蹤',3)
p('依傘面圖案、色彩或輔助標記進行原型測試，評估辨識六款雨傘及追蹤移動的可行方式。辨識方法與攝影機配置經測試後定案，需避免投影光線、遮擋及相似圖案造成誤判。')
h('(二) 雨景回饋與成果產生',3)
p('程式依雨傘辨識結果載入相應雨景，並依追蹤資訊更新互動效果。互動成果與 QR Code 應能順利連結延伸內容，且在下一次體驗前完成狀態重置。人物剪影可作為原型評估項目，其呈現方式及是否納入正式版本，依整體視覺與效能測試確認。')
h('三、設備與軟體需求',2)
p('實體雨傘規劃六款；攝影機負責影像擷取與追蹤，投影機負責雨景輸出，電腦執行辨識、動畫與成果整合。攝影機、投影機及電腦的數量與規格，依測試及展場條件確認。網頁程式與 p5.js 可作為既有原型的延伸方向，實際工具依辨識與投影整合需求選定。')
h('四、測試與驗收重點',2)
h('(一) 操作與投影穩定性',3)
p('測試六款雨傘與國家雨景的對應、撐傘與移動時的追蹤穩定性、遮擋處理、光線變化及畫面延遲。另確認不同身高與操作速度的觀眾皆能理解引導並完成體驗。')
h('(二) 延伸內容可用性',3)
p('確認互動成果清楚可讀、QR Code 可掃描、連結可開啟，且氣候知識與日常行動內容易於理解。若成果涉及觀眾影像，應於實作時明確告知使用方式並確認保存與清除機制。')

d.add_page_break()
h('伍、製作推進與預期成果',1)
h('一、製作階段',2)
h('(一) 內容與視覺開發',3)
p('確認六國名單、俚語來源與中文意譯，建立雨景素材清單；完成六款傘面設計、代表場景試作與 LOGO，並整理氣候知識及社群延伸內容。')
h('(二) 原型與系統整合',3)
p('先以單款雨傘驗證影像辨識、移動追蹤與投影回饋，再擴充至六款雨傘與六國雨景；串接互動成果、QR Code 與延伸內容，確認完整體驗流程。')
h('(三) 展場與成果測試',3)
p('調整鏡頭、投影、光線與操作動線，測試辨識穩定性、觀眾引導及成果連結。依試用回饋修正內容與操作，完成可供展出的整合版本。實際日期與個人分工由團隊排定。')
h('二、預期成果與成效',2)
h('(一) 作品交付成果',3)
p('完成六款俚語圖案雨傘、六國暴雨俚語動態雨景、影像辨識與即時投影系統，以及互動成果的 QR Code、社群延伸內容與展場操作說明。')
h('(二) 氣候意識與跨文化理解',3)
p('期望觀眾認識極端降雨對生活與環境的影響，提升對氣候變遷議題的關注；並透過各國暴雨俚語，理解不同文化的語言特色與多元表達方式。可透過現場觀察及體驗後簡短回饋，了解觀眾是否能說明俚語意象、氣候議題及一項可實踐的環境友善行動。')
h('三、後續確認事項',2)
p('團隊需確認其餘三國名單與語言查核、個人分工與時程、設備借用與規格、場地尺寸、素材及聲音授權、成果產生方式與連結維護。預算依六款雨傘製作、設備取得、輸出物及展場需求，取得實際報價後編列。')

# Two-level linked contents; chapter pages are intentionally fixed by page breaks.
page=2
for title,level,bm in headings:
 if level==1: page+=1
 if level>2: continue
 x=toc_anchor.insert_paragraph_before()
 x.paragraph_format.first_line_indent=Pt(0); x.paragraph_format.left_indent=Cm(0.6 if level==2 else 0)
 x.paragraph_format.space_after=Pt(3); x.paragraph_format.line_spacing=1.1
 x.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),WD_ALIGN_PARAGRAPH.RIGHT if False else 2,1)
 link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),bm)
 r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=title; r.append(t); link.append(r); x._p.append(link)
 x.add_run('\t'+str(page))
toc_anchor._p.getparent().remove(toc_anchor._p)
f=s.footer.paragraphs[0]; f.alignment=1
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
d.core_properties.title='俚雨畢業製作企畫書'
out=root/'俚雨_畢業製作企畫書_修訂版.docx'
d.save(out)
print(out)
