from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

root=Path(__file__).resolve().parent.parent
d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(2.3); s.bottom_margin=Cm(2.3); s.left_margin=Cm(2.5); s.right_margin=Cm(2.5)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Microsoft JhengHei'; st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(8)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.25
d.styles['Title'].font.size=Pt(30)
for st in d.styles:
 for el in list(st._element.iter(qn('w:pBdr'))): el.getparent().remove(el)
d.styles['Heading 1'].font.size=Pt(19)
d.styles['Heading 2'].font.size=Pt(13)

def p(t): return d.add_paragraph(t)
def h(t): return d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers): c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row): c.text=txt
 for i,row in enumerate(t.rows):
  for c,w in zip(row.cells,widths):
   c.width=Cm(w); tcPr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EDF0' if i==0 else 'FFFFFF'); tcPr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   tcPr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); margins.append(e)
   tcPr.append(margins)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3); para.paragraph_format.line_spacing=1.15
    for run in para.runs: run.font.size=Pt(10); run.bold=i==0
  if i==0:
   e=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(e)
 return t

p('畢業製作企畫書').alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3): p('')
a=d.add_paragraph('俚雨','Title'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
a=d.add_paragraph('同一場雨，世界會怎麼說？','Subtitle'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('以世界暴雨俚語為主題的互動投影作品').alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3): p('')
p('組員').alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['1121834　徐曼玲','1121819　游祐瑄','1121821　陳品霓','1122032　黃麗文','1122006　黃翊禎']:
 p(text).alignment=WD_ALIGN_PARAGRAPH.CENTER
p('')
p('指導老師　陳彥彰、趙伯堯').alignment=WD_ALIGN_PARAGRAPH.CENTER

page('一　企畫概要與創作背景')
h('作品概要')
p('「俚雨」以世界各地形容暴雨的俚語為創作素材，將語言中的想像轉化為動態雨景。觀眾透過旋轉雨傘切換國家，並在投影畫面中看見自己的剪影與各國俚語意象相遇。作品以文化差異作為入口，引導觀眾思考極端降雨與氣候危機的共同處境。')
h('創作動機')
p('人類曾用想像替大雨取名。英國的「下貓下狗」、葡萄牙的「下小刀」與希臘的「掉椅子腿」，以幽默、誇張或敬畏的方式描繪暴雨。這些語言意象反映不同文化理解天氣的方式，也提供將抽象語言轉化為視覺創作的可能。')
p('本作品從簡報所提出的氣候議題出發：當極端降雨愈加猛烈，過去誇張的語言隱喻是否仍足以形容眼前的雨？我們希望讓觀眾在有趣的互動體驗中，重新感受雨的力量，並思考日常語言與環境變化之間的關係。')
h('議題定位')
p('企畫對應 SDGs 第13項「氣候行動」。作品將暴雨視為跨越國界的共同議題，透過各國不同的表述，呈現文化想像與當代氣候處境的連結。')
h('創作目標')
p('蒐集各國暴雨俚語及其意象，建立可切換的動態場景。以雨傘作為操作媒介，使觀眾能直覺參與作品。將即時人物剪影融入雨景，讓觀眾從觀看者成為畫面的一部分，並留下對文化語言與氣候議題的思考。')

page('二　核心概念與文化內容')
h('概念轉譯')
p('作品依序將文化語彙轉化為雨的意象，再發展為動態視覺與互動體驗。同一個展演空間呈現不同文化的「雨」，觀眾透過操作比較各國對暴雨的想像。')
h('各國俚語與視覺素材')
table(['國家或地區','俚語','雨景意象'],[
 ['英國',"It’s raining cats and dogs.",'貓和狗'],
 ['希臘','Ρίχνει καρεκλοπόδαρα.','椅子腿'],
 ['法國','Il pleut des cordes.','繩子，可發展細繩與粗麻繩'],
 ['瑞典','Det regnar småspik.','小釘子'],
 ['葡萄牙','Está chovendo canivetes.','小刀'],
 ['南非','Dit reën ou vrouens met knopkieries.','拿著棍棒的老太太'],
 ['波蘭','Rzuca żabami.','青蛙'],
 ['巴西','Chovem cobras e lagartos.','蛇和蜥蜴'],
], [2.4,8.1,5.5])
h('內容呈現方式')
p('切換國家後，先顯示國旗，再以淡化轉場帶出該國俚語及對應雨景。俚語文字與動態物件共同建立場景辨識度，人物剪影則保持為觀眾與畫面之間的連結。')
h('語言內容確認')
p('正式展出前需確認各俚語的拼寫、語意及使用語境，並統一各國文字與中文意象說明的呈現方式。')

page('三　視覺設計與互動體驗')
h('動畫與美術方向')
p('簡報提出剪影風格，以英國場景為例探索三種方向：國旗配色搭配高對比向量印花、國旗配色搭配手繪插畫，以及暗色調情境剪影搭配倫敦街頭背景。後續可透過場景試作，確認雨景物件、背景與人物剪影的辨識度，再統一各國動畫的視覺語彙。')
h('標誌設計')
p('LOGO 以「俚雨」為作品名稱，延續語言與雨的主題。標誌需兼顧投影畫面、展場入口與作品說明的使用需求，最終版本配合整體美術方向定稿。')
h('觀眾互動流程')
table(['階段','觀眾操作與畫面回應'],[
 ['進場','牆面播放大雨，呈現「同一場雨，世界會怎麼說？」作為引導。'],
 ['操作','觀眾轉動雨傘，攝影機追蹤傘面標記並判斷旋轉。'],
 ['切換','系統切換國家，顯示國旗，國旗淡化後帶出俚語。'],
 ['體驗','投影呈現該國俚語雨景，並疊合觀眾的人物與雨傘剪影。'],
 ['再次探索','觀眾停止旋轉後，可再次操作雨傘切換至其他國家。'],
], [2.4,13.6])
h('展場配置方向')
p('以牆面作為主要投影面，前方保留觀眾持傘及轉動的操作區。攝影機一擷取人物影像，攝影機二由上往下拍攝傘面，以利追蹤顏色標記。電腦整合影像處理與互動程式，再輸出至投影機。')
p('場地尺寸、投影距離、攝影機高度與觀眾站位，需依實際展場及設備測試結果調整。配置時應避免遮擋投影光路，並保留安全的持傘活動空間。')

page('四　技術架構與設備規劃')
h('軟硬體設備')
table(['項目','數量','用途'],[
 ['攝影機','2台','分別擷取人物影像與傘面影像。'],
 ['投影機','待確認','將合成畫面投影至牆面。'],
 ['電腦','待確認','執行網頁互動程式、影像處理與畫面輸出。'],
 ['雨傘與顏色標記','待確認','提供旋轉操作與位置追蹤。'],
 ['單純背景或綠幕','依現場需求','降低背景干擾，提升人物分離穩定度。'],
 ['網頁程式與 p5.js','依開發需求','建立動畫、即時影像與互動整合。'],
], [4,2.2,9.8])
h('人物剪影處理')
p('攝影機一擷取人物影像，透過背景分離或人物去背取得遮罩（Mask）。程式將人物與雨傘輪廓填滿黑色，背景設為透明，再與雨景動畫合成並即時投影。單純背景或綠幕可降低環境干擾；背景顏色應避免與人物服裝及追蹤標記混淆。')
h('雨傘旋轉辨識')
p('攝影機二由上方拍攝傘面，以顏色標記的位置變化追蹤旋轉。依簡報設定，旋轉角度下限為180度，以避免手抖誤觸。連續多圈旋轉只觸發一次國家切換，停止旋轉後才允許再次觸發。')
h('畫面整合')
p('互動程式接收旋轉判定後切換國家，依序呈現國旗、俚語文字與雨景，並保持人物剪影同步更新。轉場需兼顧畫面連續性與辨識速度，使觀眾理解操作與回應之間的關係。')
h('主要測試項目')
p('確認人物輪廓完整度、顏色標記追蹤穩定度、180度觸發判定及停止後重置機制。另需測試快速旋轉、連續多圈、遮擋、光線改變及投影延遲，依結果調整判定與現場配置。')

page('五　製作推進與預期成果')
h('製作階段')
p('以下為建議的製作順序，實際日期與工作分配由團隊確認後填入。')
table(['階段','主要工作','完成項目'],[
 ['內容確認','整理八國俚語，確認語言與意象。','俚語內容及場景素材清單'],
 ['視覺試作','比較剪影風格，製作代表場景及 LOGO。','美術方向與動畫樣稿'],
 ['互動原型','實作人物去背、傘面追蹤與旋轉判定。','可操作的互動原型'],
 ['整合測試','串接八國雨景、文字、國旗與轉場。','完整展演程式'],
 ['展場測試','調整投影、光線、鏡頭與觀眾站位。','展場配置及穩定運行版本'],
], [2.8,7.3,5.9])
h('預期成果')
p('完成以「俚雨」為題的互動投影作品，包含八國暴雨俚語雨景、雨傘旋轉切換功能、即時人物剪影與國旗及文字轉場。展場提供清楚的操作引導與作品說明，使觀眾能透過體驗理解不同文化的雨之意象，並連結氣候行動議題。')
h('成果驗收方向')
p('觀眾能依引導完成操作並切換國家；俚語文字、國旗與雨景對應正確；人物剪影可辨識且與雨景同步；連續旋轉不會重複誤觸，停止後可重新操作。整體作品需在展場光線與設備條件下穩定呈現。')
h('後續確認事項')
p('團隊需確認製作時程與個人分工、設備規格與借用方式、展場尺寸、最終美術風格、LOGO 定稿，以及素材與聲音的使用安排。預算依設備取得方式與實際報價編列，金額另行確認。')

d.core_properties.title='俚雨畢業製作企畫書'
d.core_properties.subject='互動投影畢業製作'
d.save(root/'俚雨_畢業製作企畫書.docx')
print(root/'俚雨_畢業製作企畫書.docx')

